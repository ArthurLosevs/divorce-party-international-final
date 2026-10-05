from pathlib import Path
import json, hashlib, csv
from pypdf import PdfReader
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent
out = root / 'evidence-extracted'
out.mkdir(exist_ok=True)
inventory = []
for path in sorted(root.iterdir()):
    if path.suffix.lower() not in ['.pdf', '.xlsx', '.csv']:
        continue
    item = {'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'bytes': path.stat().st_size}
    if path.suffix == '.pdf':
        reader = PdfReader(path)
        pages = [{'page': i+1, 'text': p.extract_text()} for i,p in enumerate(reader.pages)]
        item['pages'] = len(pages)
        (out / (path.stem + '.txt')).write_text('\n\n'.join(f'PAGE {p["page"]}\n{p["text"]}' for p in pages), encoding='utf-8')
    elif path.suffix == '.xlsx':
        wb = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
        sheets = []
        for ws in wb:
            rows = []
            for row in ws:
                cells = [{'cell': c.coordinate, 'value': c.value, 'cached': cached[ws.title][c.coordinate].value, 'comment': c.comment.text if c.comment else ''} for c in row if c.value is not None or c.comment]
                if cells: rows.append(cells)
            sheets.append({'sheet': ws.title, 'state': ws.sheet_state, 'rows': rows})
        item['sheets'] = [s['sheet'] for s in sheets]
        (out / (path.stem + '.json')).write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        (out / (path.stem + '.txt')).write_text('\n\n'.join('SHEET '+s['sheet']+'\n'+'\n'.join(' | '.join(f'{c["cell"]}: {c["value"]}'+(f' [cached {c["cached"]}]' if str(c['value']).startswith('=') else '')+(f' [comment: {c["comment"]}]' if c['comment'] else '') for c in r) for r in s['rows']) for s in sheets), encoding='utf-8')
    else:
        item['rows'] = len(list(csv.DictReader(path.open(encoding='utf-8-sig'))))
    inventory.append(item)
(out/'manifest.json').write_text(json.dumps(inventory, indent=2), encoding='utf-8')
print(json.dumps(inventory, indent=2))
